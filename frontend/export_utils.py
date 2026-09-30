from io import BytesIO
from pathlib import Path
import re
from typing import Optional

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Image,
    Table,
    TableStyle,
)
from reportlab.lib import colors


def sanitize_text(text: str) -> str:
    replacements = {
        "\u2018": "'",
        "\u2019": "'",
        "\u201c": '"',
        "\u201d": '"',
        "\u2013": "-",
        "\u2014": "-",
        "\u00a0": " ",
    }
    for old, new in replacements.items():
        text = text.replace(old, new)
    return re.sub(r"[^\x09\x0A\x0D\x20-\x7E]", "", text)


def _split_lines(text: str) -> list[str]:
    return [line.strip() for line in sanitize_text(text).splitlines() if line.strip()]


def format_txt(text: str) -> bytes:
    return sanitize_text(text).encode("utf-8")


def format_docx(
    text: str,
    doc_type: str,
    logo_bytes: Optional[bytes] = None,
) -> bytes:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    if logo_bytes:
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run()
        run.add_picture(BytesIO(logo_bytes), width=Inches(1.1))

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(doc_type.upper())
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(16)

    for line in _split_lines(text):
        p = document.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.15
        r = p.add_run(line)
        r.font.name = "Times New Roman"
        r.font.size = Pt(11)
        if re.match(r"^(\d+[\.\)]|[A-Z][A-Z\s&-]{3,})", line):
            r.bold = True

    terms = []
    for line in _split_lines(text):
        if line.startswith(("-", "•", "*")):
            terms.append(line.lstrip("-•* ").strip())

    if terms:
        document.add_paragraph()
        heading = document.add_paragraph()
        heading.add_run("Key Terms").bold = True
        table = document.add_table(rows=1, cols=2)
        table.style = "Table Grid"
        table.rows[0].cells[0].text = "No."
        table.rows[0].cells[1].text = "Term"
        for i, term in enumerate(terms, 1):
            cells = table.add_row().cells
            cells[0].text = str(i)
            cells[1].text = term

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Generated with LegalEase — AI-assisted draft. Review before use.")

    output = BytesIO()
    document.save(output)
    return output.getvalue()


def format_pdf(
    text: str,
    doc_type: str,
    logo_bytes: Optional[bytes] = None,
) -> bytes:
    output = BytesIO()
    doc = SimpleDocTemplate(
        output,
        pagesize=A4,
        rightMargin=50,
        leftMargin=50,
        topMargin=55,
        bottomMargin=55,
        title=doc_type,
        author="LegalEase",
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "LegalTitle",
        parent=styles["Title"],
        fontName="Helvetica-Bold",
        fontSize=15,
        leading=19,
        alignment=TA_CENTER,
        spaceAfter=14,
    )
    body_style = ParagraphStyle(
        "LegalBody",
        parent=styles["BodyText"],
        fontName="Times-Roman",
        fontSize=10.5,
        leading=15,
        alignment=TA_LEFT,
        spaceAfter=7,
    )
    heading_style = ParagraphStyle(
        "LegalHeading",
        parent=body_style,
        fontName="Times-Bold",
        spaceBefore=8,
        spaceAfter=5,
    )

    story = []
    if logo_bytes:
        logo_path = BytesIO(logo_bytes)
        img = Image(logo_path, width=75, height=75, kind="proportional")
        img.hAlign = "CENTER"
        story.append(img)
        story.append(Spacer(1, 5))

    story.append(Paragraph(doc_type.upper(), title_style))

    for line in _split_lines(text):
        escaped = (
            line.replace("&", "&amp;")
            .replace("<", "&lt;")
            .replace(">", "&gt;")
        )
        if re.match(r"^(\d+[\.\)]|[A-Z][A-Z\s&-]{3,})", line):
            story.append(Paragraph(escaped, heading_style))
        elif line.startswith(("-", "•", "*")):
            story.append(Paragraph("• " + escaped.lstrip("-•* "), body_style))
        else:
            story.append(Paragraph(escaped, body_style))

    def footer(canvas, document):
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(colors.grey)
        canvas.drawCentredString(
            A4[0] / 2,
            28,
            "LegalEase — AI-assisted draft. Review before use.",
        )
        canvas.restoreState()

    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return output.getvalue()


def suggested_filename(doc_type: str) -> str:
    clean = re.sub(r"[^A-Za-z0-9]+", "_", doc_type.strip()).strip("_")
    return (clean or "LegalEase_Document")[:80]
